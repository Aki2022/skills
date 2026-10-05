#!/usr/bin/env python3
"""export_speaker_notes.py — pptx のスピーカーノートを Markdown に書き出す。

⑤で `own-project-update` の `attach-document --notes-file` へ渡す素材を作る（SPEC-document-publish）。
ノートのあるスライドだけを `## スライド N` 見出しつきで書く。ノートが 1 枚も無い pptx と、読めない
ファイルは非 0 で止める（空の公開素材を黙って作らない。④で全スライドのノート有無は検査済みのはず）。

使い方:
  python3 export_speaker_notes.py <成果物>.pptx                    # 標準出力
  python3 export_speaker_notes.py <成果物>.pptx --output process/speaker_notes.md
"""

import argparse
import re
import sys
from pathlib import Path

try:
    from pptx import Presentation
except ImportError:  # pragma: no cover - 環境の欠落
    raise SystemExit("python-pptx is required: pip install python-pptx")


# PowerPoint の Shift+Enter は python-pptx の text で垂直タブ(\v)になる。str.splitlines() はこれらで行を分けるので、
# vault の document note に入れる前に普通の改行へ揃える（own-project-update も取り込み時に同じ正規化をする）。
_LINE_SEPARATORS = re.compile("[\x0b\x0c\x1c\x1d\x1e\x85\u2028\u2029]")


def export(pptx_path: Path) -> str:
    try:
        deck = Presentation(str(pptx_path))
    except Exception as exc:  # 壊れた zip・非 pptx など。原因は原文のまま出す
        raise SystemExit(f"cannot read {pptx_path.name}: {exc}")
    blocks = []
    for number, slide in enumerate(deck.slides, start=1):
        if not slide.has_notes_slide:
            continue
        text = _LINE_SEPARATORS.sub("\n", slide.notes_slide.notes_text_frame.text).strip()
        if text:
            blocks.append(f"## スライド {number}\n\n{text}\n")
    if not blocks:
        raise SystemExit(f"{pptx_path.name}: no speaker notes found on any slide")
    return "\n".join(blocks)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("pptx", type=Path)
    parser.add_argument("--output", type=Path, help="書き出し先（既定: 標準出力）")
    args = parser.parse_args()
    if not args.pptx.is_file():
        raise SystemExit(f"not a file: {args.pptx.name}")
    text = export(args.pptx)
    if args.output:
        args.output.write_text(text, encoding="utf-8")
    else:
        sys.stdout.write(text)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

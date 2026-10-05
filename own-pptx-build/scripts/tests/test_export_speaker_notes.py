"""export_speaker_notes.py: ⑤の attach-document へ渡すスピーカーノートの書き出し。

pptx の slide.notes_slide から、ノートのあるスライドだけを `## スライド N` 見出しつきで書き出す。
ノートが 1 枚も無い・読めないファイルは黙って空出力にせず、非 0 で止める（空の公開素材を作らない）。
"""

import importlib.util
import subprocess
import sys
from pathlib import Path

import pytest
from pptx import Presentation

SCRIPT = Path(__file__).resolve().parents[1] / "export_speaker_notes.py"


def make_deck(path: Path, notes: list[str | None]) -> Path:
    prs = Presentation()
    layout = prs.slide_layouts[6]
    for text in notes:
        slide = prs.slides.add_slide(layout)
        if text is not None:
            slide.notes_slide.notes_text_frame.text = text
    prs.save(path)
    return path


def run(*args: str) -> subprocess.CompletedProcess:
    return subprocess.run([sys.executable, str(SCRIPT), *args], capture_output=True, text=True)


def test_exports_only_slides_that_have_notes_with_their_numbers(tmp_path: Path) -> None:
    deck = make_deck(tmp_path / "d.pptx", ["最初の語り。", None, "三枚目の語り。\n二行目。"])
    result = run(str(deck))
    assert result.returncode == 0, result.stderr
    assert result.stdout == "## スライド 1\n\n最初の語り。\n\n## スライド 3\n\n三枚目の語り。\n二行目。\n"


def test_deck_without_any_notes_is_an_error_not_an_empty_success(tmp_path: Path) -> None:
    deck = make_deck(tmp_path / "d.pptx", [None, ""])
    result = run(str(deck))
    assert result.returncode != 0
    assert result.stdout == ""
    assert "notes" in result.stderr


def test_missing_or_non_pptx_file_is_an_error(tmp_path: Path) -> None:
    assert run(str(tmp_path / "nope.pptx")).returncode != 0
    bogus = tmp_path / "bogus.pptx"
    bogus.write_text("not a zip", encoding="utf-8")
    assert run(str(bogus)).returncode != 0


def test_output_option_writes_the_file(tmp_path: Path) -> None:
    deck = make_deck(tmp_path / "d.pptx", ["語り。"])
    out = tmp_path / "speaker_notes.md"
    assert run(str(deck), "--output", str(out)).returncode == 0
    assert out.read_text(encoding="utf-8") == "## スライド 1\n\n語り。\n"


def test_soft_line_breaks_become_plain_newlines(tmp_path: Path) -> None:
    """Shift+Enter は python-pptx で \\v（垂直タブ）になる。vault へ渡す前に普通の改行へ直す。"""
    deck = make_deck(tmp_path / "d.pptx", ["一行目\x0b二行目"])
    result = run(str(deck))
    assert result.returncode == 0, result.stderr
    assert "\x0b" not in result.stdout
    assert result.stdout == "## スライド 1\n\n一行目\n二行目\n"
